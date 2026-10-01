"""El filtro de Año que llega a Inicio, Resumen, Localidades y el Excel.

`resumen_localidad` y `resumen_ccte` son tablas materializadas por clave
(ccte, provincia, localidad) SIN columna de año: hasta que el año entró por
`listar_resumen_localidad`, poner el Año en el panel no cambiaba nada en
esas salidas -- `/localities`, `/top-localities`, `/ccte-summary` y el Excel
devolvían exactamente los mismos renglones con y sin filtro, mientras el
resto del dashboard sí se acotaba.

El caso que el bug tapaba es una localidad medida en DOS años, así que el
fixture la arma así: `Salta Capital` con 3 mediciones de 2025 y 3 de 2026
(6 en la tabla materializada) y `Cachi` sólo con 2025.
"""
import io

import pandas as pd
from openpyxl import load_workbook


def _excel_bytes(anio: str) -> bytes:
    """Un .xlsx con el formato de ENACOM (8 filas de metadata antes del
    header, igual que los archivos reales) y las fechas en el año pedido."""
    df = pd.DataFrame({
        "Resultado": ["1,5", "3.2", "10"],
        "Fecha": [f"20/03/{anio}", f"20/03/{anio}", f"21/03/{anio}"],
        "Hora": ["10:00:00", "10:30:00", "09:00:00 a.m."],
        "Lat": [-34.6037, -34.6040, -34.6050],
        "Lon": [-58.3816, -58.3820, -58.3830],
        "Sonda": ["S1", "S1", "S2"],
    })
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        pd.DataFrame([[f"meta{i}"] for i in range(8)]).to_excel(
            writer, index=False, header=False, startrow=0)
        df.to_excel(writer, index=False, startrow=8)
    buf.seek(0)
    return buf.read()


def _importar(client, localidad: str, anio: str, expediente: str = "EXP-1") -> int:
    """Carga 3 mediciones de `anio` en una localidad del CCTE Salta y
    devuelve los registros nuevos."""
    resp = client.post(
        "/api/import",
        data={"ccte": "Salta", "provincia": "Salta", "localidad": localidad,
              "expediente": expediente},
        files={"archivos": (f"{anio}.xlsx", _excel_bytes(anio),
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200, resp.text
    return resp.json()["registros_nuevos"]


def _dos_anios(client):
    assert _importar(client, "Salta Capital", "2025") == 3
    assert _importar(client, "Salta Capital", "2026") == 3
    assert _importar(client, "Cachi", "2025", expediente="EXP-2") == 3


def _por_localidad(filas) -> dict:
    return {f["localidad"]: f for f in filas}


def _resumen_ccte(client, **params) -> dict:
    return {f["ccte"]: f for f in client.get("/api/ccte-summary", params=params).json()}


# --- /localities ------------------------------------------------------------

def test_localities_sin_anio_muestra_toda_la_historia(client):
    """La tabla materializada sigue siendo la puerta cuando no hay año:
    resume los dos años juntos."""
    _dos_anios(client)

    por_localidad = _por_localidad(client.get("/api/localities").json())

    assert set(por_localidad) == {"Salta Capital", "Cachi"}
    assert por_localidad["Salta Capital"]["mediciones"] == 6
    assert por_localidad["Cachi"]["mediciones"] == 3


def test_localities_con_anio_recalcula_en_vivo(client):
    """Con el año puesto la respuesta cambia de verdad: mismas columnas,
    filas acotadas a ese año y fechas que arrancan en ese año."""
    _dos_anios(client)

    c2025 = _por_localidad(client.get("/api/localities", params={"anio": 2025}).json())
    c2026 = _por_localidad(client.get("/api/localities", params={"anio": 2026}).json())

    # 2025 tiene las dos localidades, 3 mediciones cada una.
    assert set(c2025) == {"Salta Capital", "Cachi"}
    assert {k: v["mediciones"] for k, v in c2025.items()} == {"Salta Capital": 3, "Cachi": 3}
    # 2026 sólo la que se midió ese año: si el filtro no actuara, esta
    # localidad aparecería igual con 6 mediciones.
    assert set(c2026) == {"Salta Capital"}
    assert c2026["Salta Capital"]["mediciones"] == 3
    assert c2026["Salta Capital"]["fecha_inicio"].startswith("2026-")

    # Un año sin datos no devuelve la tabla entera: devuelve nada.
    assert client.get("/api/localities", params={"anio": 1999}).json() == []


# --- /top-localities --------------------------------------------------------

def test_top_localities_con_anio_acota_el_ranking(client):
    _dos_anios(client)

    todos = client.get("/api/top-localities", params={"metric": "mediciones"}).json()
    assert [(t["localidad"], t["valor"]) for t in todos] == [("Salta Capital", 6), ("Cachi", 3)]

    solo_2026 = client.get("/api/top-localities",
                           params={"metric": "mediciones", "anio": 2026}).json()
    assert [(t["localidad"], t["valor"]) for t in solo_2026] == [("Salta Capital", 3)]

    c2025 = client.get("/api/top-localities", params={"metric": "mediciones", "anio": 2025}).json()
    assert {t["localidad"] for t in c2025} == {"Salta Capital", "Cachi"}
    assert all(t["valor"] == 3 for t in c2025)


# --- /ccte-summary ----------------------------------------------------------

def test_ccte_summary_con_anio_suma_solo_ese_anio(client):
    _dos_anios(client)

    completo = _resumen_ccte(client)
    assert len(completo) == 7, "los 7 CCTE fijos tienen que seguir ahí"
    assert completo["Salta"]["mediciones"] == 9
    assert completo["Salta"]["localidades"] == 2

    c2025 = _resumen_ccte(client, anio=2025)
    assert len(c2025) == 7, "el año no saca CCTE de la lista, sólo los pone en 0"
    assert c2025["Salta"]["mediciones"] == 6
    assert c2025["Salta"]["localidades"] == 2

    c2026 = _resumen_ccte(client, anio=2026)
    assert c2026["Salta"]["mediciones"] == 3
    assert c2026["Salta"]["localidades"] == 1
    assert all(v["mediciones"] == 0 for k, v in c2026.items() if k != "Salta")


# --- Excel ------------------------------------------------------------------

def test_excel_de_resumen_baja_con_el_anio(client):
    _dos_anios(client)

    resp = client.get("/api/reports/excel", params={"anio": 2025})
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")

    ws = load_workbook(io.BytesIO(resp.content)).active
    filas = [[c.value for c in fila] for fila in ws.iter_rows()]
    encabezado = next(f for f in filas if f and "CCTE" in f)
    i_loc = encabezado.index("Localidad")
    i_med = encabezado.index("Mediciones")
    datos = {f[i_loc]: f[i_med] for f in filas if f and f[i_loc] in ("Salta Capital", "Cachi")}

    assert datos == {"Salta Capital": 3, "Cachi": 3}, "el Excel tiene que bajar filtrado"

    # Sin filas el endpoint responde 404, igual que con cualquier otro filtro
    # que no coincide (misma rama que `?ccte=Inexistente`).
    assert client.get("/api/reports/excel", params={"anio": 1999}).status_code == 404
