"""Exports del Lote 13: el Excel de Resumen, las tres opciones (Excel, Word y
PDF) por localidad y por CCTE, y el PDF del mapa.

Lo que se cubre aca es lo que el usuario reporta como "no funciona": no basta
con que el endpoint devuelva 200, hay que abrir el archivo. Por eso la mayoria
de los tests lee el resultado con python-docx u openpyxl en vez de mirar solo
el codigo de estado -- un Word vacio tambien responde 200.
"""
import base64
import io

import pandas as pd
from openpyxl import load_workbook
from PIL import Image

from app.db.repositories import resumen_repo
from app.services import import_service, reports as reports_service

# Columnas de la vista Resumen, escritas a mano a proposito (no se toman de
# reports_service.COLUMNAS_RESUMEN): si alguien cambia la constante sin
# cambiar la vista, el test tiene que quejarse, no repetirle la respuesta.
COLUMNAS_VISTA = ["CCTE", "Provincia", "Localidad", "Expediente(s)", "Mediciones",
                  "Máx. V/m", "Nivel", "Inicio", "Fin"]

FORMATO_VM = "0.###"      # hasta 3 decimales: los que el instrumento tiene
FORMATO_PCT = "0.####"    # hasta 4: los que imprime la vista


def _df_dos_dias() -> pd.DataFrame:
    """4 mediciones en 2 dias: 20/03 10:00-10:30 (1800 s) y
    21/03 09:00-09:20 (1200 s). El mismo fixture que usa test_centro_operativo,
    copiado para que este archivo se pueda leer sin importar otro test."""
    return pd.DataFrame({
        "Resultado": ["1,5", "3.2", "10", "2.1"],
        "Fecha": ["20/03/2025", "20/03/2025", "21/03/2025", "21/03/2025"],
        "Hora": ["10:00:00", "10:30:00", "09:00:00 a.m.", "09:20:00 a.m."],
        "Lat": [-34.6037, -34.6040, -34.6050, -34.6060],
        "Lon": [-58.3816, -58.3820, -58.3830, -58.3840],
        "Sonda": ["S1", "S1", "S2", "S2"],
    })


def _importar(conn, ccte, provincia, localidad, expediente=None, desplazamiento=0.0):
    """Una localidad con 4 mediciones. `desplazamiento` corre la latitud para
    que dos localidades del mismo CCTE no caigan en el mismo punto."""
    df = _df_dos_dias()
    df["Lat"] = df["Lat"] + desplazamiento
    return import_service.importar_lote(
        conn, ccte=ccte, provincia=provincia, localidad=localidad,
        expediente=expediente, archivos=[("a.xlsx", df)],
    )


def _abrir_excel(buf):
    return load_workbook(io.BytesIO(buf.getvalue())).active


def _textos(ws):
    return [str(c.value) for fila in ws.iter_rows() for c in fila if c.value is not None]


def _encabezado(ws):
    """Fila de encabezados de la tabla, sin contar los titulos del documento."""
    for fila in ws.iter_rows():
        vals = [c.value for c in fila]
        if "CCTE" in vals and "Localidad" in vals:
            return vals
    return None


def _filas_datos(ws):
    """Filas de la tabla: son las unicas con un entero en la columna Mediciones
    (5). Los titulos, el encabezado y el pie no lo tienen."""
    return [f for f in ws.iter_rows() if isinstance(f[4].value, int)]


# --- alcance del informe ----------------------------------------------------

def test_informe_ccte_es_la_union_de_sus_localidades(conn):
    """El informe de un CCTE baja por la misma consulta y los mismos calculos
    que el de cada localidad (los dos pasan por _construir_datos), asi que sus
    totales tienen que cuadrar exacto con la suma de las localidades."""
    _importar(conn, "Cordoba", "Cordoba", "Cordoba Capital")
    _importar(conn, "Cordoba", "Cordoba", "Villa Allende", desplazamiento=1.0)

    ccte = reports_service.obtener_datos_informe_ccte(conn, "Cordoba")
    loc1 = reports_service.obtener_datos_informe(conn, "Cordoba", "Cordoba", "Cordoba Capital")
    loc2 = reports_service.obtener_datos_informe(conn, "Cordoba", "Cordoba", "Villa Allende")

    assert ccte["total_puntos"] == 8
    assert ccte["total_puntos"] == loc1["total_puntos"] + loc2["total_puntos"]
    assert ccte["resultado_max_vm"] == max(loc1["resultado_max_vm"], loc2["resultado_max_vm"])
    assert ccte["sondas"] == ["S1", "S2"]


def test_informe_ccte_sin_datos_no_inventa_nada(conn):
    datos = reports_service.obtener_datos_informe_ccte(conn, "no existe")
    assert datos["filas"] == []
    assert datos["df"].empty
    assert "total_puntos" not in datos   # el llamador tiene que ver que no hay


# --- logo en los tres formatos ---------------------------------------------

def test_word_lleva_el_logo_ademas_del_grafico(conn):
    """Antes el Word no traia logo: el informe se veia igual que cualquier
    archivo suelto. Ahora son dos imagenes incrustadas (logo + grafico)."""
    from docx import Document

    _importar(conn, "Cordoba", "Cordoba", "Cordoba Capital")
    datos = reports_service.obtener_datos_informe(conn, "Cordoba", "Cordoba", "Cordoba Capital")

    buf = reports_service.generar_word(datos, ambito="Localidad", etiqueta="Cordoba Capital")
    doc = Document(io.BytesIO(buf.getvalue()))

    assert len(doc.inline_shapes) >= 2, "falta el logo o el grafico"
    parrafos = [p.text for p in doc.paragraphs]
    assert "\u00c1mbito del informe: Localidad" in parrafos
    assert "Localidad: Cordoba Capital" in parrafos


def test_pdf_trae_tiempo_y_sondas_que_antes_solo_ponia_el_word(conn, monkeypatch):
    """El PDF era mas pobre que el Word: no llevaba el tiempo trabajado ni las
    sondas. Se espiaron los Paragraph que llegan a reportlab, porque buscar el
    texto dentro del PDF comprimido no se puede sin otra dependencia."""
    _importar(conn, "Cordoba", "Cordoba", "Cordoba Capital")
    datos = reports_service.obtener_datos_informe(conn, "Cordoba", "Cordoba", "Cordoba Capital")

    textos, imagenes = [], []
    parrafo_original = reports_service.Paragraph
    imagen_original = reports_service.RLImage

    def espia_parrafo(texto, estilo, *args, **kwargs):
        textos.append(str(texto))
        return parrafo_original(texto, estilo, *args, **kwargs)

    def espia_imagen(*args, **kwargs):
        imagenes.append(args)
        return imagen_original(*args, **kwargs)

    monkeypatch.setattr(reports_service, "Paragraph", espia_parrafo)
    monkeypatch.setattr(reports_service, "RLImage", espia_imagen)

    buf = reports_service.generar_pdf(datos, ambito="Localidad", etiqueta="Cordoba Capital")

    assert buf.getvalue().startswith(b"%PDF")
    assert any("Tiempo total estimado de medici" in t for t in textos), "falta el tiempo"
    assert any("Sondas utilizadas" in t for t in textos), "faltan las sondas"
    assert any(str(reports_service.LOGO) == str(args[0]) for args in imagenes), \
        "el logo no llego al PDF"


# --- Excel de Resumen -------------------------------------------------------

def test_excel_resumen_replica_la_vista_con_expediente(conn):
    """El boton de Resumen baja la tabla tal cual se ve: mismas columnas que
    frontend/src/views/ResumenView.vue, incluida Expediente(s)."""
    _importar(conn, "Cordoba", "Cordoba", "Cordoba Capital", expediente="EXP-77")
    filas = resumen_repo.listar_resumen_localidad(conn)

    buf = reports_service.generar_excel_resumen(filas, filtros="Sin filtros")
    assert buf.getvalue().startswith(b"PK")

    ws = _abrir_excel(buf)
    assert _encabezado(ws) == COLUMNAS_VISTA, "las columnas se desincronizaron de la vista"

    filas_datos = _filas_datos(ws)
    assert len(filas_datos) == 1
    assert filas_datos[0][3].value == "EXP-77"   # la columna Expediente(s)

    assert len(ws._images) >= 1, "falta el logo"

    textos = _textos(ws)
    assert any("Filtros aplicados:" in t for t in textos), "no declara los filtros"
    assert any("Fecha de generaci" in t for t in textos), "no pone fecha de generacion"
    assert any("Direcci" in t and "Fiscalizaci" in t for t in textos), "falta el pie"

    # Los decimales: el valor queda CRUDO en la celda y el formato decide
    # cuanto se ve, igual que hace el frontend (3 y 4 decimales, sin rellenar
    # con ceros). Quien copie la celda se lleva el numero exacto de la base.
    assert filas_datos[0][5].number_format == FORMATO_VM
    assert filas_datos[0][6].number_format == FORMATO_PCT
    # El valor va CRUDO en la celda: el formato decide cuantos decimales se
    # ven, no redondea lo que la base guarda. El maximo de las 4 mediciones
    # sinteticas es 10, y la base lo trae como entero.
    assert float(filas_datos[0][5].value) == 10.0


def test_excel_resumen_sin_filas_aun_se_arma(conn):
    """La ruta HTTP responde 404 cuando no hay resultados; el generador por si
    solo tiene que poder armar el documento igual, para no fallar a mitad."""
    filas = resumen_repo.listar_resumen_localidad(conn)
    assert filas == []
    buf = reports_service.generar_excel_resumen(filas, filtros="Sin filtros")
    assert buf.getvalue().startswith(b"PK")


# --- Excel de informe -------------------------------------------------------

def test_excel_informe_trae_kpis_grafico_y_tablas(conn):
    """El Excel de una localidad o de un CCTE lleva los mismos bloques que el
    Word: KPIs, grafico, desglose mensual y resumen por expediente."""
    _importar(conn, "Cordoba", "Cordoba", "Cordoba Capital", expediente="EXP-9")
    datos = reports_service.obtener_datos_informe(conn, "Cordoba", "Cordoba", "Cordoba Capital")

    buf = reports_service.generar_excel_informe(datos, ambito="Localidad",
                                                etiqueta="Cordoba Capital")
    assert buf.getvalue().startswith(b"PK")

    ws = _abrir_excel(buf)
    textos = _textos(ws)
    for buscado in ("Resultado m\u00e1ximo registrado", "Tiempo total de medici\u00f3n",
                    "Desglose mensual", "Resumen por expediente",
                    "Localidad: Cordoba Capital"):
        assert any(buscado in t for t in textos), "falta '%s'" % buscado

    # logo + grafico de localidades por provincia
    assert len(ws._images) >= 2, "falta el logo o el grafico"
    assert any("EXP-9" in t for t in textos), "falta el expediente en la tabla"


def test_excel_de_ccte_identifica_el_centro(conn):
    _importar(conn, "Cordoba", "Cordoba", "Cordoba Capital")
    datos = reports_service.obtener_datos_informe_ccte(conn, "Cordoba")

    buf = reports_service.generar_excel_informe(datos, ambito="CCTE", etiqueta="Cordoba")
    ws = _abrir_excel(buf)
    textos = _textos(ws)
    assert any("CCTE: Cordoba" in t for t in textos)
    assert any("\u00c1mbito del informe: CCTE" in t for t in textos)


# --- PDF del mapa -----------------------------------------------------------

def test_pdf_del_mapa_se_arma_en_memoria(tmp_path, monkeypatch):
    """La captura del mapa pasa por reportlab para poder imprimirla en A4.
    Igual que el resto de los informes, no escribe nada en disco."""
    buf_png = io.BytesIO()
    Image.new("RGB", (800, 500), (240, 240, 240)).save(buf_png, format="PNG")

    workdir = tmp_path / "workdir"
    workdir.mkdir()
    monkeypatch.chdir(workdir)

    buf = reports_service.generar_pdf_mapa(
        buf_png.getvalue(), titulo="Mapa de mediciones RNI",
        notas=["Con filtros globales: CCTE Cordoba"],
    )

    assert buf.getvalue().startswith(b"%PDF")
    assert list(workdir.iterdir()) == []


# --- por HTTP ---------------------------------------------------------------

def _xlsx_enacom() -> bytes:
    """8 filas de metadata antes del header, igual que un archivo real."""
    df = pd.DataFrame({
        "Resultado": ["1,5", "3.2", "10"],
        "Fecha": ["20/03/2025", "20/03/2025", "21/03/2025"],
        "Hora": ["10:00:00", "10:30:00", "09:00:00 a.m."],
        "Lat": [-34.6037, -34.6040, -34.6050],
        "Lon": [-58.3816, -58.3820, -58.3830],
        "Sonda": ["S1", "S1", "S2"],
    })
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        pd.DataFrame([[f"meta{i}"] for i in range(8)]).to_excel(
            writer, index=False, header=False, startrow=0)
        df.to_excel(writer, index=False, startrow=8)
    buf.seek(0)
    return buf.read()


def _cargar(client, localidad="CABA"):
    resp = client.post(
        "/api/import",
        data={"ccte": "Buenos Aires", "provincia": "Buenos Aires",
              "localidad": localidad, "expediente": "EXP-1"},
        files={"archivos": ("muestra.xlsx", _xlsx_enacom(),
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200, resp.text


LOC = {"ccte": "Buenos Aires", "provincia": "Buenos Aires", "localidad": "CABA",
       "ambito": "Localidad"}
CCTE = {"ccte": "Buenos Aires"}


def _descarga_ok(client, ruta, params, mago, nombre):
    resp = client.get(ruta, params=params)
    assert resp.status_code == 200, "%s -> %s %s" % (ruta, resp.status_code, resp.text[:300])
    assert resp.content.startswith(mago), "%s: el archivo no abre" % ruta
    assert 'attachment; filename="%s"' % nombre in resp.headers["content-disposition"], ruta
    return resp


def test_las_tres_opciones_de_localidad_descargan_por_http(client):
    _cargar(client)
    for ruta, mago, nombre in (
        ("/api/reports/excel/localidad", b"PK", "Informe_RNI_CABA.xlsx"),
        ("/api/reports/word", b"PK", "Informe_RNI_CABA.docx"),
        ("/api/reports/pdf", b"%PDF", "Informe_RNI_CABA.pdf"),
    ):
        _descarga_ok(client, ruta, LOC, mago, nombre)


def test_las_tres_opciones_de_ccte_descargan_por_http(client):
    """Las opciones del Centro operativo: mismos formatos, otro alcance."""
    _cargar(client)
    for ruta, mago, nombre in (
        ("/api/reports/excel/ccte", b"PK", "Informe_RNI_Buenos Aires.xlsx"),
        ("/api/reports/word/ccte", b"PK", "Informe_RNI_Buenos Aires.docx"),
        ("/api/reports/pdf/ccte", b"%PDF", "Informe_RNI_Buenos Aires.pdf"),
    ):
        _descarga_ok(client, ruta, CCTE, mago, nombre)


def test_excel_de_resumen_por_http_respeta_el_filtro(client):
    _cargar(client)
    _cargar(client, localidad="La Plata")

    resp = client.get("/api/reports/excel")
    assert resp.status_code == 200
    ws = _abrir_excel(io.BytesIO(resp.content))
    assert len(_filas_datos(ws)) == 2

    resp = client.get("/api/reports/excel", params={"localidad": "CABA"})
    assert resp.status_code == 200
    ws = _abrir_excel(io.BytesIO(resp.content))
    filas = _filas_datos(ws)
    assert len(filas) == 1
    assert filas[0][2].value == "CABA"
    assert any("Localidad: CABA" in str(f[0].value) for f in ws.iter_rows() if f[0].value)


def test_informe_inexistente_responde_404_y_no_un_archivo_vacio(client):
    """Antes devolvia un Word con 'Total de puntos medidos: 0': un 200 que
    parece funcionar y no dice nada."""
    for ruta in ("/api/reports/word", "/api/reports/pdf", "/api/reports/excel/localidad"):
        resp = client.get(ruta, params={"ccte": "x", "provincia": "y", "localidad": "no existe"})
        assert resp.status_code == 404, ruta

    resp = client.get("/api/reports/word/ccte", params=CCTE)
    assert resp.status_code == 404
    resp = client.get("/api/reports/excel", params={"ccte": "no existe"})
    assert resp.status_code == 404


def test_pdf_del_mapa_por_http(client):
    buf_png = io.BytesIO()
    Image.new("RGB", (800, 500), (240, 240, 240)).save(buf_png, format="PNG")

    resp = client.post("/api/reports/map-pdf", json={
        "png": base64.b64encode(buf_png.getvalue()).decode(),
        "titulo": "Mapa de mediciones RNI",
        "notas": ["Con filtros globales: CCTE Buenos Aires"],
    })
    assert resp.status_code == 200, resp.text
    assert resp.content.startswith(b"%PDF")
    assert 'attachment; filename="Mapa_RNI.pdf"' in resp.headers["content-disposition"]

    # no es un 500 mudo: lo que no es imagen se rechaza con un mensaje
    resp = client.post("/api/reports/map-pdf", json={"png": "esto no es base64!!!"})
    assert resp.status_code == 400, resp.status_code

    # base64 valido pero de algo que no es PNG
    resp = client.post("/api/reports/map-pdf", json={
        "png": base64.b64encode(b"texto plano, no un PNG").decode()})
    assert resp.status_code == 400, resp.status_code
