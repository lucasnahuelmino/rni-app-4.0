"""Generación de informes Word/PDF de una localidad.

Port funcional de sections/export_informes.py del sistema Streamlit actual,
con una corrección de la Auditoría Fase 1 (hallazgo A9): el informe Word
original se escribía a un archivo físico en el directorio de trabajo del
proceso antes de leerlo de vuelta para la descarga -- acá todo se arma en
memoria (BytesIO), nunca se toca el disco, así que no hay riesgo de
colisión entre usuarios concurrentes ni archivos huérfanos en el servidor.

El gráfico usa matplotlib en vez de Plotly+Kaleido (el sistema actual
depende de Kaleido solo para poder exportar el PNG del informe) -- reduce
una dependencia pesada para un gráfico simple de barras. Si el equipo
prefiere mantener el estilo visual exacto de Plotly, se puede reemplazar
sin tocar el resto del pipeline.
"""
from __future__ import annotations

import io
import sqlite3
from datetime import datetime, timezone

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Image as RLImage
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from app.calculations.dates import calcular_tiempo_trabajado_segundos, format_timedelta_long


def _armar_grafico_localidades_por_provincia_ccte(filas: list[dict]) -> io.BytesIO | None:
    df = pd.DataFrame(filas)
    if df.empty or not {"provincia", "ccte", "localidad"}.issubset(df.columns):
        return None

    resumen = df.groupby(["provincia", "ccte"])["localidad"].nunique().reset_index(name="cantidad")
    if resumen.empty:
        return None

    fig, ax = plt.subplots(figsize=(6, 3.5))
    for ccte_val, grupo in resumen.groupby("ccte"):
        ax.bar(grupo["provincia"], grupo["cantidad"], label=ccte_val)
    ax.set_ylabel("Localidades")
    ax.legend(fontsize=7)
    fig.tight_layout()

    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150)
    plt.close(fig)
    buf.seek(0)
    return buf


def obtener_datos_informe(conn: sqlite3.Connection, ccte: str, provincia: str, localidad: str) -> dict:
    filas = [dict(r) for r in conn.execute(
        "SELECT * FROM mediciones WHERE ccte = ? AND provincia = ? AND localidad = ?",
        (ccte, provincia, localidad),
    ).fetchall()]

    df = pd.DataFrame(filas)
    if df.empty:
        return {"filas": [], "df": df}

    df["fecha_hora"] = pd.to_datetime(df["fecha_hora"], errors="coerce")

    idx_max = df["resultado_vm"].idxmax() if df["resultado_vm"].notna().any() else None
    fila_max = df.loc[idx_max] if idx_max is not None else None

    fecha_min = df["fecha_hora"].min() if df["fecha_hora"].notna().any() else None
    fecha_max = df["fecha_hora"].max() if df["fecha_hora"].notna().any() else None

    tiempo_trabajado_seg = calcular_tiempo_trabajado_segundos(df)

    sondas = sorted(df["sonda"].dropna().astype(str).unique().tolist()) if "sonda" in df else []

    resumen_mensual = pd.DataFrame()
    if df["fecha_hora"].notna().any():
        df_mes = df.dropna(subset=["fecha_hora"]).copy()
        df_mes["mes"] = df_mes["fecha_hora"].dt.to_period("M").astype(str)
        resumen_mensual = df_mes.groupby("mes").agg(
            puntos=("resultado_vm", "count"),
            fecha_inicio=("fecha_hora", "min"),
            fecha_fin=("fecha_hora", "max"),
        ).reset_index()
        horas_por_mes = []
        for mes, grupo in df_mes.groupby("mes"):
            horas_por_mes.append({"mes": mes, "horas": format_timedelta_long(calcular_tiempo_trabajado_segundos(grupo))})
        resumen_mensual = resumen_mensual.merge(pd.DataFrame(horas_por_mes), on="mes")

    expedientes_df = pd.DataFrame()
    if "expediente" in df.columns and df["expediente"].notna().any():
        expedientes_df = df.groupby("expediente").agg(
            puntos=("resultado_vm", "count"),
            max_vm=("resultado_vm", "max"),
            ccte=("ccte", lambda x: ", ".join(sorted(x.dropna().unique()))),
            provincias=("provincia", lambda x: ", ".join(sorted(x.dropna().unique()))),
            localidades=("localidad", lambda x: ", ".join(sorted(x.dropna().unique()))),
        ).reset_index().sort_values("max_vm", ascending=False)

    grafico = _armar_grafico_localidades_por_provincia_ccte(filas)

    return {
        "filas": filas,
        "total_puntos": len(df),
        "resultado_max_vm": None if fila_max is None else float(fila_max["resultado_vm"]),
        "resultado_max_pct": None if fila_max is None else fila_max.get("resultado_pct"),
        "localidad_max": None if fila_max is None else fila_max.get("localidad"),
        "provincia_max": None if fila_max is None else fila_max.get("provincia"),
        "ccte_max": None if fila_max is None else fila_max.get("ccte"),
        "fecha_hora_max": None if fila_max is None else fila_max.get("fecha_hora"),
        "fecha_min": fecha_min, "fecha_max": fecha_max,
        "tiempo_trabajado_seg": tiempo_trabajado_seg,
        "sondas": sondas,
        "resumen_mensual": resumen_mensual,
        "expedientes_df": expedientes_df,
        "grafico": grafico,
    }


def generar_word(datos: dict, *, localidad: str, ambito: str) -> io.BytesIO:
    doc = Document()
    doc.add_heading("Informe de Mediciones RNI", level=1)
    doc.add_paragraph(f"Ámbito del informe: {ambito}")
    doc.add_paragraph(f"Localidad: {localidad}")
    doc.add_paragraph(f"Fecha de generación: {datetime.now(timezone.utc).strftime('%d/%m/%Y %H:%M:%S')} UTC")
    doc.add_paragraph(f"Total de puntos medidos: {datos.get('total_puntos', 0)}")

    if datos.get("resultado_max_vm") is not None:
        pct = datos.get("resultado_max_pct")
        texto = f"Resultado máximo registrado: {datos['resultado_max_vm']:.2f} V/m"
        if pct is not None:
            texto += f" ({pct:.2f} % del límite)"
        doc.add_paragraph(texto)
        doc.add_paragraph(
            f"Ubicación del máximo: {datos.get('localidad_max')}, {datos.get('provincia_max')} "
            f"(CCTE {datos.get('ccte_max')})"
        )

    if datos.get("tiempo_trabajado_seg"):
        doc.add_paragraph(f"Tiempo total estimado de medición: {format_timedelta_long(datos['tiempo_trabajado_seg'])}")

    if datos.get("sondas"):
        doc.add_paragraph(f"Sondas utilizadas: {', '.join(datos['sondas'])}")

    if datos.get("grafico") is not None:
        doc.add_heading("Localidades por provincia y CCTE", level=2)
        doc.add_picture(datos["grafico"], width=None)

    resumen_mensual = datos.get("resumen_mensual")
    if resumen_mensual is not None and not resumen_mensual.empty:
        doc.add_heading("Desglose mensual", level=2)
        table = doc.add_table(rows=1, cols=5)
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text, hdr[4].text = (
            "Mes", "Puntos", "Horas trabajadas", "Fecha inicio", "Fecha fin",
        )
        for _, row in resumen_mensual.iterrows():
            cells = table.add_row().cells
            cells[0].text = str(row["mes"])
            cells[1].text = str(row["puntos"])
            cells[2].text = row["horas"]
            cells[3].text = row["fecha_inicio"].strftime("%d/%m/%Y %H:%M") if pd.notna(row["fecha_inicio"]) else "-"
            cells[4].text = row["fecha_fin"].strftime("%d/%m/%Y %H:%M") if pd.notna(row["fecha_fin"]) else "-"

    expedientes_df = datos.get("expedientes_df")
    if expedientes_df is not None and not expedientes_df.empty:
        doc.add_heading("Resumen por expediente", level=2)
        table = doc.add_table(rows=1, cols=6)
        hdr = table.rows[0].cells
        for i, titulo in enumerate(["Expediente", "Puntos", "Max (V/m)", "CCTE", "Provincias", "Localidades"]):
            hdr[i].text = titulo
        for _, row in expedientes_df.iterrows():
            cells = table.add_row().cells
            cells[0].text = str(row["expediente"])
            cells[1].text = str(row["puntos"])
            cells[2].text = f"{row['max_vm']:.2f}" if pd.notna(row["max_vm"]) else "-"
            cells[3].text = str(row["ccte"])
            cells[4].text = str(row["provincias"])
            cells[5].text = str(row["localidades"])

    buffer = io.BytesIO()
    doc.save(buffer)  # se guarda directo en memoria, nunca en disco
    buffer.seek(0)
    return buffer


def generar_pdf(datos: dict, *, localidad: str, ambito: str) -> io.BytesIO:
    buffer = io.BytesIO()
    pdf = SimpleDocTemplate(buffer, pagesize=A4)
    styles = getSampleStyleSheet()
    story = [
        Paragraph("Informe de Mediciones RNI", styles["Title"]),
        Spacer(1, 6),
        Paragraph(f"Ámbito del informe: {ambito}", styles["Heading2"]),
        Spacer(1, 12),
        Paragraph(f"<b>Localidad:</b> {localidad}", styles["Normal"]),
        Paragraph(f"<b>Fecha de generación:</b> {datetime.now(timezone.utc).strftime('%d/%m/%Y %H:%M:%S')} UTC",
                  styles["Normal"]),
        Paragraph(f"<b>Total de puntos medidos:</b> {datos.get('total_puntos', 0)}", styles["Normal"]),
    ]

    if datos.get("resultado_max_vm") is not None:
        pct = datos.get("resultado_max_pct")
        texto = f"<b>Resultado máximo registrado:</b> {datos['resultado_max_vm']:.2f} V/m"
        if pct is not None:
            texto += f" ({pct:.2f} % del límite)"
        story.append(Paragraph(texto, styles["Normal"]))
        story.append(Paragraph(
            f"<b>Ubicación del máximo:</b> {datos.get('localidad_max')}, {datos.get('provincia_max')} "
            f"(CCTE {datos.get('ccte_max')})", styles["Normal"],
        ))

    story.append(Spacer(1, 16))

    if datos.get("grafico") is not None:
        story.append(RLImage(datos["grafico"], width=400, height=250))
        story.append(Spacer(1, 16))

    resumen_mensual = datos.get("resumen_mensual")
    if resumen_mensual is not None and not resumen_mensual.empty:
        story.append(Paragraph("<b>Desglose por mes</b>", styles["Heading2"]))
        for _, row in resumen_mensual.iterrows():
            story.append(Paragraph(
                f"Mes {row['mes']}: {row['puntos']} puntos, horas trabajadas: {row['horas']}.", styles["Normal"],
            ))
        story.append(Spacer(1, 12))

    expedientes_df = datos.get("expedientes_df")
    if expedientes_df is not None and not expedientes_df.empty:
        story.append(Paragraph("<b>Resumen por expediente</b>", styles["Heading2"]))
        for _, row in expedientes_df.iterrows():
            story.append(Paragraph(
                f"Expediente {row['expediente']}: {row['puntos']} puntos, "
                f"máx {row['max_vm']:.2f} V/m, CCTE: {row['ccte']}.", styles["Normal"],
            ))

    pdf.build(story)
    buffer.seek(0)
    return buffer
